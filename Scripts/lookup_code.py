"""
Reads Master_List/departments.xlsx and returns the record for a doc code:
which standard applies, which template to use, and its control category.
Claude Code calls this first, before drafting anything.
"""
import sys
import json
from pathlib import Path
import openpyxl

MASTER_LIST_PATH = Path(__file__).resolve().parent.parent / "Master_List" / "departments.xlsx"


def _load_sheet():
    wb = openpyxl.load_workbook(MASTER_LIST_PATH, data_only=True)
    return wb["MasterMatrix"] if "MasterMatrix" in wb.sheetnames else wb.active


def get_doc_entry(doc_code: str) -> dict:
    ws = _load_sheet()
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers) if h}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx["Doc Code"]] == doc_code:
            return {
                "doc_code": row[idx["Doc Code"]],
                "doc_name": row[idx["Doc Name"]],
                "standard_id": row[idx["Standard ID"]],
                "control_category": row[idx["Control Category"]],
                "template_file": row[idx["Template File"]],
                "source_input_type": row[idx["Source Input Type"]],
            }
    raise ValueError(f'Doc code "{doc_code}" not found in {MASTER_LIST_PATH.name}')


def list_all_doc_codes() -> list:
    ws = _load_sheet()
    headers = [c.value for c in ws[1]]
    code_idx = headers.index("Doc Code")
    return [row[code_idx] for row in ws.iter_rows(min_row=2, values_only=True) if row[code_idx]]


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python lookup_code.py <DOC_CODE>")
        sys.exit(1)
    try:
        print(json.dumps(get_doc_entry(sys.argv[1]), indent=2))
    except ValueError as e:
        print(str(e))
        print("Known codes:", ", ".join(list_all_doc_codes()) or "(none yet)")
        sys.exit(1)
